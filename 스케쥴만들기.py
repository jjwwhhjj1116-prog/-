# -*- coding: utf-8 -*-
"""
컨코스트 마감팀 프로젝트 스케쥴 + 직원 업무달력 v6.0
- 완전 실시간 수식 연동 (고전 함수만: MATCH/INDEX/COUNTIF/ISNA — 모든 엑셀 호환)
- 색상 없음(회색), 행 높이 전부 30
실행(최초 1회 파일 생성): py 스케쥴만들기.py
"""
import os
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from datetime import timedelta

HERE=os.path.dirname(os.path.abspath(__file__))
OUT=os.path.join(HERE,"컨코스트_마감팀_스케쥴.xlsx")
N_ROWS=60; K_SLOTS=10; YEAR=2026
BA_N=N_ROWS*9; BA_END=BA_N+1

DOMESTIC=[("최영배 본부장","최영배"),("조한빈 실장","조한빈"),("김재헌 수석","김재헌"),
          ("양한규 수석","양한규"),("성대용 수석","성대용"),("원종수 수석","원종수"),
          ("송영길 수석","송영길"),("이은지 책임","이은지"),("송치영 책임","송치영"),
          ("남은주 책임","남은주"),("임승주 선임","임승주"),("임창열 선임","임창열"),("김수겸 프로","김수겸")]
VIET=[("VIET 내부1팀","VIET 내부1팀"),("VIET 내부2팀","VIET 내부2팀"),("VIET 내부3팀","VIET 내부3팀"),
      ("VIET 외부팀","VIET 외부팀"),("VIET 창호팀","VIET 창호팀"),("VIET 조적팀","VIET 조적팀")]
FIN_NAMES=[m for _,m in DOMESTIC]+[m for _,m in VIET]
OTHER=["신동헌","박용진","장범선","오승균","박가림","이성희","이정철","박소현","김채원"]
DISC_NAMES=["X"]+[m for _,m in DOMESTIC]+[m for _,m in VIET]+OTHER
STATUS=["시작전","진행중","검토","완료","보류"]; PRIO=["긴급","높음","보통","낮음"]
GONG=["PM","검토","내역","가설","세대","내부","외부","창호","조적"]
DEF_DUR={"PM":0,"검토":3,"내역":3,"가설":3,"세대":5,"내부":7,"외부":5,"창호":4,"조적":4}

# ── 무채색 ──
HDR="595959"; FLD="D9D9D9"; BAR="BFBFBF"; TXT="000000"
WHITE_B=Font(color="FFFFFF",bold=True,size=10)
THIN=Side(style="thin",color="C8C8C8"); BORDER=Border(THIN,THIN,THIN,THIN)
CEN=Alignment(horizontal="center",vertical="center",wrap_text=True)
LEFT=Alignment(horizontal="left",vertical="center",wrap_text=True)
def fill(h): return PatternFill("solid",fgColor=h)

GROUPS={"base":"기본정보","disc":"공종 담당","fin":"마감 세부담당","sch":"일정 / 상태","meet":"회의일정","iss":"이슈사항"}
FIELDS=[("순번",6,"base",None),("No",10,"base",None),("업무명",34,"base",None),("거래처",16,"base",None),
        ("연면적",9,"base",None),("타입수",7,"base",None),("선행작업",12,"base",None),
        ("마감",9,"disc","disc"),("구조",9,"disc","disc"),("토목",9,"disc","disc"),
        ("PM",8,"fin","fin"),("검토",8,"fin","fin"),("내역",8,"fin","fin"),("가설",8,"fin","fin"),
        ("세대",8,"fin","fin"),("내부",8,"fin","fin"),("외부",8,"fin","fin"),("창호",8,"fin","fin"),("조적",8,"fin","fin"),
        ("수주일",11,"sch",None),("작업시작일",11,"sch",None),("납품일",11,"sch",None),
        ("소요일수\n(자동)",8,"sch",None),("우선순위",9,"sch","prio"),("진행률\n(%)",8,"sch",None),
        ("현재 상태",10,"sch","status"),("비고",24,"sch",None),
        ("인원",8,"meet",None),("일자",11,"meet",None),("내용",20,"meet",None),
        ("마감",9,"iss",None),("구조",9,"iss",None),("토목",9,"iss",None)]
CL={i:get_column_letter(2+i) for i in range(len(FIELDS))}
C_NM=CL[2]; C_SD=CL[20]; C_ND=CL[21]; C_STT=CL[25]
GONG_COL={g:CL[10+i] for i,g in enumerate(GONG)}

def _d(s): y,m,d=map(int,s.split("-")); return date(y,m,d)
SEED=[
 ("2026010","구리 수택E구역 주택재개발 2,4BL 견적","대형건설",95000,4,"","송치영","신동헌","X",
  {"PM":"송치영","조적":"VIET 조적팀"},"2026-03-06","2026-03-20","보통",60,"진행중",""),
 ("2026011","광양바이오매스 발전소 EPC 설계변경","엔지니어링",22000,1,"","최영배","X","오승균",
  {"PM":"최영배"},"2026-03-11","2026-03-20","높음",80,"진행중","설계변경분"),
 ("2026012","평택고덕 A-64BL 통합형 공공주택","종합건설",64000,3,"도면접수","조한빈","박용진","X",
  {"PM":"조한빈","내부":"임승주","창호":"VIET 창호팀","조적":"VIET 조적팀"},"2026-03-11","2026-04-03","긴급",30,"진행중","부대·공용부"),
 ("2026013","남산스퀘어 주거복합","시행사",48000,2,"","이은지","X","X",
  {"PM":"이은지","세대":"이은지","외부":"임창열","창호":"김수겸","조적":"김수겸"},"2026-06-15","2026-07-03","보통",10,"시작전","김수겸 창호+조적"),
 ("2026014","마이다스아이티 사옥","건설사",18000,1,"","조한빈","X","X",
  {"PM":"조한빈","내부":"임승주","창호":"김수겸"},"2026-06-03","2026-06-19","높음",50,"진행중",""),
 ("2026015","위례복정역세권 공동주택","공공",72000,3,"","성대용","장범선","X",
  {"PM":"성대용","가설":"원종수"},"2026-05-25","2026-06-12","보통",100,"완료","검토완료(달력제외)"),
]

def set_heights(ws,rows,h=30):
    for r in rows: ws.row_dimensions[r].height=h

# ===================== 진행중 =====================
def build_jinhaeng(wb):
    ws=wb.create_sheet("진행중")
    ws["A1"]="마감팀 프로젝트 일정 — 담당자만 입력하면 [업무 달력]이 직원별로 실시간 자동 정리됩니다"
    ws["A1"].font=Font(bold=True,size=12,color=TXT); ws.column_dimensions["A"].width=2
    i=0
    while i<len(FIELDS):
        g=FIELDS[i][2]; j=i
        while j+1<len(FIELDS) and FIELDS[j+1][2]==g: j+=1
        c=ws[f"{CL[i]}3"]; c.value=GROUPS[g]; c.fill=fill(HDR); c.font=WHITE_B; c.alignment=CEN; c.border=BORDER
        if j>i: ws.merge_cells(f"{CL[i]}3:{CL[j]}3")
        i=j+1
    for idx,(name,w,g,dv) in enumerate(FIELDS):
        cell=ws[f"{CL[idx]}4"]; cell.value=name; cell.fill=fill(FLD)
        cell.font=Font(bold=True,size=9,color=TXT); cell.alignment=CEN; cell.border=BORDER
        ws.column_dimensions[CL[idx]].width=w
    first,last=5,5+N_ROWS-1
    for r in range(first,last+1):
        ws[f"{CL[0]}{r}"]=f'=IF({C_NM}{r}="","",ROW()-4)'
        ws[f"{CL[22]}{r}"]=f'=IF(AND({C_SD}{r}<>"",{C_ND}{r}<>""),NETWORKDAYS({C_SD}{r},{C_ND}{r}),"")'
        for idx,(name,w,g,dv) in enumerate(FIELDS):
            cc=ws[f"{CL[idx]}{r}"]; cc.border=BORDER
            cc.alignment=LEFT if name in ("업무명","거래처","선행작업","비고","내용") else CEN
        for cl in (CL[19],CL[20],CL[21],CL[28]): ws[f"{cl}{r}"].number_format="yyyy-mm-dd"
        ws[f"{CL[24]}{r}"].number_format='0"%"'
    for si,row in enumerate(SEED):
        r=first+si
        no_,nm,gc_,area,typ,pre,mg,st,tm,roles,sd,nd,pr,pg,stt,bg=row
        base={CL[1]:no_,CL[2]:nm,CL[3]:gc_,CL[4]:area,CL[5]:typ,CL[6]:pre,CL[7]:mg,CL[8]:st,CL[9]:tm,
              CL[20]:_d(sd),CL[21]:_d(nd),CL[23]:pr,CL[24]:pg,CL[25]:stt,CL[26]:bg}
        for cl,v in base.items():
            if v not in (None,""): ws[f"{cl}{r}"]=v
        for g,who in roles.items(): ws[f"{GONG_COL[g]}{r}"]=who
    def dv(lst,letters):
        d=DataValidation(type="list",formula1='"%s"'%",".join(lst),allow_blank=True); ws.add_data_validation(d)
        for col in letters: d.add(f"{col}{first}:{col}{last}")
    dv(DISC_NAMES,[CL[i] for i,f in enumerate(FIELDS) if f[3]=="disc"])
    dv(FIN_NAMES,[CL[i] for i,f in enumerate(FIELDS) if f[3]=="fin"])
    dv(PRIO,[CL[23]]); dv(STATUS,[CL[25]])
    # 상태 회색만 (색상 없음)
    ws.conditional_formatting.add(f"{CL[0]}{first}:{CL[26]}{last}",
        Rule(type="expression",dxf=DifferentialStyle(fill=PatternFill("solid",bgColor="EDEDED"),font=Font(color="808080")),
             formula=[f'${C_STT}{first}="완료"']))
    ws.freeze_panes=f"{CL[3]}5"
    set_heights(ws,[3,4]+list(range(first,last+1)),30)
    return ws

# ===================== 설정 =====================
def build_settings(wb):
    ws=wb.create_sheet("설정")
    ws["A1"]="공종별 표준 투입일수(근무일). PM=0=전기간. 숫자만 바꾸면 달력 막대 길이 자동"
    ws["A1"].font=Font(bold=True,size=11,color=TXT)
    for j,h in enumerate(["공종","표준일수"],start=1):
        ws.cell(3,j,h).fill=fill(HDR); ws.cell(3,j).font=WHITE_B; ws.cell(3,j).alignment=CEN; ws.cell(3,j).border=BORDER
    for i,g in enumerate(GONG,start=4):
        ws[f"A{i}"]=g; ws[f"B{i}"]=DEF_DUR[g]
        ws[f"A{i}"].border=BORDER; ws[f"B{i}"].border=BORDER; ws[f"B{i}"].alignment=CEN
    ws.column_dimensions["A"].width=10; ws.column_dimensions["B"].width=10
    set_heights(ws,[3]+list(range(4,4+len(GONG))),30)
    return ws

# ===================== 배정 (숨김) =====================
def build_baejeong(wb):
    ws=wb.create_sheet("배정"); ws.sheet_state="hidden"
    for j,h in enumerate(["담당","업무명","공종","시작","종료","순번","키"],start=1): ws.cell(1,j,h)
    DUR='설정!$A$4:$B$12'
    r=2
    for p in range(N_ROWS):
        pr=5+p
        for g in GONG:
            gc=GONG_COL[g]
            ws.cell(r,1).value=(f'=IF(AND(진행중!{C_NM}{pr}<>"",진행중!{C_SD}{pr}<>"",진행중!{gc}{pr}<>"",'
                                f'진행중!{gc}{pr}<>"X",진행중!{C_STT}{pr}<>"완료"),진행중!{gc}{pr},"")')
            ws.cell(r,2).value=f'=IF($A{r}="","",진행중!{C_NM}{pr})'
            ws.cell(r,3).value=g
            ws.cell(r,4).value=f'=IF($A{r}="","",진행중!{C_SD}{pr})'
            ws.cell(r,5).value=(f'=IF($A{r}="","",IF(VLOOKUP($C{r},{DUR},2,0)=0,진행중!{C_ND}{pr},'
                                f'MIN(WORKDAY(진행중!{C_SD}{pr},VLOOKUP($C{r},{DUR},2,0)-1),진행중!{C_ND}{pr})))')
            ws.cell(r,6).value=f'=IF($A{r}="","",COUNTIF($A$2:$A{r},$A{r}))'   # 담당별 순번
            ws.cell(r,7).value=f'=IF($A{r}="","",$A{r}&"|"&$F{r})'             # 키
            ws.cell(r,4).number_format="yyyy-mm-dd"; ws.cell(r,5).number_format="yyyy-mm-dd"
            r+=1
    return ws

# ===================== 업무 달력 =====================
def build_calendar(wb):
    ws=wb.create_sheet("📅 업무 달력")
    BB=f'배정!$B$2:$B${BA_END}'; BC=f'배정!$C$2:$C${BA_END}'; BD=f'배정!$D$2:$D${BA_END}'
    BE=f'배정!$E$2:$E${BA_END}'; BG=f'배정!$G$2:$G${BA_END}'
    heads=[("A","직원",14),("B","업무 (공종)",26),("C","시작",10),("D","종료",10)]
    for col,name,w in heads:
        ws[f"{col}2"]=name; ws[f"{col}2"].fill=fill(HDR); ws[f"{col}2"].font=WHITE_B
        ws[f"{col}2"].alignment=CEN; ws[f"{col}2"].border=BORDER; ws.column_dimensions[col].width=w
    for col in ("E","F","G"): ws.column_dimensions[col].hidden=True
    ws["A1"]="직원 업무 달력 (마감팀) — 진행중 입력 즉시 자동 반영 · 막대=공종 표준 투입일수"
    ws["A1"].font=Font(bold=True,size=12,color=TXT)
    d0=date(YEAR,1,1); d1=date(YEAR,12,31); ndays=(d1-d0).days+1; SC=8
    mcols={}
    for i in range(ndays):
        d=d0+timedelta(days=i); col=SC+i; Lc=get_column_letter(col)
        ws.column_dimensions[Lc].width=2.4
        hc=ws.cell(2,col); hc.value=d; hc.number_format="d"
        hc.font=Font(size=7,color=TXT); hc.fill=fill(FLD); hc.alignment=CEN; hc.border=BORDER
        mcols.setdefault((d.year,d.month),[]).append(col)
    for (y,m),cs in mcols.items():
        a,b=min(cs),max(cs); cc=ws.cell(1,a,f"{m}월"); cc.fill=fill(HDR); cc.font=WHITE_B; cc.alignment=CEN
        if b>a: ws.merge_cells(start_row=1,start_column=a,end_row=1,end_column=b)
    last_col=SC+ndays-1; LASTL=get_column_letter(last_col)
    ws.cell(1,1,f"{YEAR}년").font=Font(bold=True,size=11,color=TXT)
    def band(r,text):
        ws.cell(r,1,text).font=Font(bold=True,size=10,color=TXT)
        for c in range(1,8): ws.cell(r,c).fill=fill(FLD); ws.cell(r,c).border=BORDER
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=7); ws.cell(r,1).alignment=LEFT
    r=3; heightrows=[1,2]
    for sect,roster in [("[ 국내 마감팀 ]",DOMESTIC),("[ VIET QS (베트남 마감팀) ]",VIET)]:
        band(r,sect); heightrows.append(r); r+=1
        for disp,match in roster:
            r0=r
            for k in range(1,K_SLOTS+1):
                ws[f"E{r}"]=match
                ws[f"F{r}"]=(f'=IF(ISNA(MATCH($E{r}&"|"&{k},{BG},0)),"",MATCH($E{r}&"|"&{k},{BG},0))')
                ws[f"B{r}"]=f'=IF($F{r}="","",INDEX({BB},$F{r})&" ("&INDEX({BC},$F{r})&")")'
                ws[f"C{r}"]=f'=IF($F{r}="","",INDEX({BD},$F{r}))'
                ws[f"D{r}"]=f'=IF($F{r}="","",INDEX({BE},$F{r}))'
                ws[f"G{r}"]=f'=IF($F{r}="","",INDEX({BC},$F{r}))'
                ws[f"C{r}"].number_format="mm-dd"; ws[f"D{r}"].number_format="mm-dd"
                ws[f"B{r}"].alignment=LEFT; ws[f"B{r}"].font=Font(size=9)
                for c in range(1,5): ws.cell(r,c).border=BORDER
                heightrows.append(r); r+=1
            ws.cell(r0,1,disp).font=Font(bold=True,size=9); ws.cell(r0,1).alignment=CEN
            for rr in range(r0,r): ws.cell(rr,1).border=BORDER
            ws.merge_cells(start_row=r0,start_column=1,end_row=r-1,end_column=1)
    last_row=r-1
    # 막대: 회색 1색
    ws.conditional_formatting.add(f"H3:{LASTL}{last_row}",
        Rule(type="expression",dxf=DifferentialStyle(fill=PatternFill("solid",bgColor=BAR)),
             formula=['AND($C3<>"",H$2>=$C3,H$2<=$D3)']))
    ws.freeze_panes="H3"; ws.sheet_view.showGridLines=False
    set_heights(ws,heightrows,30)
    return ws

def build_guide(wb):
    ws=wb.create_sheet("작성 가이드")
    ws["B1"]="사용법"; ws["B1"].font=Font(bold=True,size=12,color=TXT)
    rows=[("입력","[진행중] 시트에만 입력. 마감/구조/토목 + PM~조적 칸에 담당자 선택, 작업시작일·납품일 입력"),
          ("자동연동","입력 즉시 [업무 달력]에 직원별로 막대 자동 생성. 새로고침 불필요"),
          ("공종 일수","[설정] 시트 숫자로 공종별 투입일수 조정(PM=0=전기간)"),
          ("줄 분리","한 직원이 동시에 여러 공종이면 여러 줄에 표시(직원당 최대 %d개)"%K_SLOTS),
          ("완료 제외","현재상태=완료면 달력에서 자동 제외"),
          ("주의","날짜는 반드시 날짜형식(2026-03-06). 배정 시트(숨김)는 계산용이니 건드리지 마세요")]
    for i,(a,b) in enumerate(rows,start=3):
        ws[f"B{i}"]=a; ws[f"C{i}"]=b; ws[f"B{i}"].font=Font(bold=True)
        ws[f"B{i}"].alignment=Alignment(vertical="top",wrap_text=True); ws[f"C{i}"].alignment=Alignment(vertical="top",wrap_text=True)
    ws.column_dimensions["B"].width=14; ws.column_dimensions["C"].width=78

def main():
    wb=Workbook(); wb.remove(wb.active)
    try:
        wb.calculation.calcMode="auto"        # 자동 계산 강제
        wb.calculation.fullCalcOnLoad=True    # 열 때 전체 재계산
        wb.calculation.forceFullCalc=True
    except Exception: pass
    build_jinhaeng(wb); build_settings(wb); build_baejeong(wb); build_calendar(wb); build_guide(wb)
    order=["진행중","📅 업무 달력","설정","작성 가이드","배정"]
    wb._sheets.sort(key=lambda s: order.index(s.title) if s.title in order else 99)
    wb.save(OUT); print("[done]",OUT)

if __name__=="__main__":
    main()
